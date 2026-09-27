import docx

doc = docx.Document('SEF GROUP TEMPLATE.docx')
print('=== Paragraphs ===')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_info = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.font.bold) for r in p.runs]
        print(f'P{i}: text="{p.text}", style={p.style.name}, align={p.alignment}, runs={runs_info}')

print('\n=== Tables ===')
for t_idx, t in enumerate(doc.tables):
    print(f'Table {t_idx}:')
    for r_idx, row in enumerate(t.rows):
        cells_info = []
        for cell in row.cells:
            for p in cell.paragraphs:
                p_runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.font.bold) for r in p.runs]
                cells_info.append(p_runs)
        print(f'  Row {r_idx}: {cells_info}')

print('\n=== Styles in Document ===')
for s in doc.styles:
    if s.type == docx.enum.style.WD_STYLE_TYPE.PARAGRAPH and ('Heading' in s.name or 'Normal' in s.name):
        font = s.font
        print(f'Style: {s.name}, font={font.name}, size={font.size.pt if font.size else None}')
